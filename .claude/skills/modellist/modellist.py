#!/usr/bin/env python3
"""modellist —— 模型信息落盘 model.xlsx（MODEL sheet，模板感知）。

两个落盘时机（--stage）:
  arch : YAML 生成时落架构列（MODEL_TAG/INPUT/VFE/.../NMS + FLOPs/PARAMs）
  perf : 训练结束出 eval 后落性能列（2D*/3D* mAP）
  all  : 两者（默认）

模板规则（铁律）:
  - 只写值，不改样式：字体 Times New Roman / 居中 / 列宽 / 表头 / 空列分隔均不触碰
  - 列位置从表头动态读取（禁止硬编码列号）
  - 新增行复制模板数据行（第 2 行）的单元格样式
"""
import argparse
import importlib.util
import json
import re
import sys
from copy import copy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]  # .claude/skills/modellist -> 工程根
sys.path.insert(0, str(ROOT))

ARCH_COLS = ['MODEL_TAG', 'INPUT', 'VFE', '3DBACKBONE', '2DBACKBONE', 'NECK', 'HEAD',
             'NMS', 'FLOPs', 'PARAMs']
PERF_COLS = ['2DmAP', '2DCYCmAP', '2DPERmAP', '2DCARmAP', '2DTRUCKmAP', '2DOBmAP',
             '3DmAP', '3DCYCmAP', '3DPERmAP', '3DCARmAP', '3DTRUCKmAP', '3DOBmAP']
# eval 类名 -> mAP 列后缀（OB=障碍物 type7，当前不参与训练/评估）
CLASS2COL = {'Cyclist': 'CYC', 'Pedestrian': 'PER', 'Car': 'CAR', 'Truck': 'TRUCK'}
METRIC_CLASSES = ['CYC', 'PER', 'CAR', 'TRUCK']


def _log(msg, level='INFO'):
    print(f'[modellist][{level}] {msg}', file=sys.stderr)


def _fmt(v):
    """列表 -> '[a,b,c]'（无空格）。"""
    return '[%s]' % ','.join(str(x) for x in v)


# -- cfg 读取（pcdet 口径，含 _BASE_CONFIG_ 递归合并）-------------------------
def load_cfg(yaml_path: Path):
    from easydict import EasyDict
    from pcdet.config import cfg_from_yaml_file
    cfg = EasyDict()
    cfg_from_yaml_file(str(yaml_path), cfg)
    return cfg


# -- 架构列命名（规则与 mynote/McRadarPillar.md「模型表命名规则」一致）--------
def derive_arch(cfg) -> dict:
    m = cfg.MODEL
    feats = cfg.DATA_CONFIG.POINT_FEATURE_ENCODING.used_feature_list
    cols = {'INPUT': '[%s]' % ','.join(feats)}

    vfe = m.get('VFE') or {}
    nf = vfe.get('NUM_FILTERS')
    cols['VFE'] = ('PN' + _fmt(nf)) if (nf and 'VFE' in vfe.get('NAME', '')) \
        else vfe.get('NAME', '—')

    bb3 = m.get('BACKBONE_3D') or {}
    if 'PillarAttention' in bb3.get('NAME', ''):
        cols['3DBACKBONE'] = 'PAtt[head:%d]' % bb3.NUM_HEADS
    else:
        cols['3DBACKBONE'] = bb3.get('NAME', '—')

    bb2 = m.get('BACKBONE_2D') or {}
    name2 = bb2.get('NAME', '')
    if 'BaseBEVBackbone' in name2:
        cols['2DBACKBONE'] = 'PP' + _fmt(bb2.LAYER_NUMS) + '*' + _fmt(bb2.NUM_FILTERS)
    elif 'REPDWC' in name2.upper():
        cols['2DBACKBONE'] = 'RepDwc' + _fmt(bb2.LAYER_NUMS) + '*' + _fmt(bb2.OUT_CHANNELS)
    else:
        cols['2DBACKBONE'] = name2 or '—'

    # NECK: 独立 neck 模块优先；否则 2D backbone 内置 deblock 上采样+concat
    #   写 UPSAMPLE[UPSAMPLE_STRIDES]+CONCAT[C=各级上采样通道和]
    neck = m.get('NECK')
    if neck and neck.get('NAME'):
        cols['NECK'] = neck['NAME'] + (_fmt([neck['CHANNELS']]) if neck.get('CHANNELS') else '')
    else:
        ups = bb2.get('NUM_UPSAMPLE_FILTERS')
        strides = bb2.get('UPSAMPLE_STRIDES')
        if ups and strides:
            cols['NECK'] = 'UPSAMPLE' + _fmt(strides) + '+CONCAT' + _fmt([sum(ups)])
        else:
            cols['NECK'] = ('CONCAT' + _fmt([sum(ups)])) if ups else '—'

    dh = m.get('DENSE_HEAD') or {}
    cols['HEAD'] = 'anchorfree' if 'Center' in dh.get('NAME', '') else 'anchorbased'

    nms = (m.get('POST_PROCESSING') or {}).get('NMS_CONFIG') or dh.get('NMS_CONFIG') or {}
    if nms.get('NMS_TYPE'):
        parts = [nms.NMS_TYPE.replace('nms_', ''), nms.get('NMS_THRESH')]
        if nms.get('MULTI_CLASSES_NMS'):
            parts.append('mc')
        cols['NMS'] = 'NMS[%s]' % ','.join(str(x) for x in parts)
    else:
        cols['NMS'] = '—'
    return cols


# -- run 目录定位 + params/FLOPs/mAP 提取 -------------------------------------
def _has_artifacts(d: Path) -> bool:
    return any((d / x).exists() for x in
               ('best.pth', 'model_store.yaml', 'eval', 'ckpt'))


def find_run_dir(yaml_stem: str, override: str | None) -> Path | None:
    """run 目录匹配规则（output/train_log/* 全域扫）:
    1. --run_dir 显式指定
    2. 目录名以 _<yaml名> 结尾（如 *_msr_centerhead2d）
    3. baseline 特例: yaml 名 = <ds>_radarpillar 而 tag=<ds>（如 *_msr）
    多候选取有产物且 mtime 最新的（排除空壳中断 run）。
    """
    if override:
        p = Path(override)
        return p if p.exists() else None
    runs_root = ROOT / 'output' / 'train_log'
    if not runs_root.exists():
        return None
    cands = [d for d in runs_root.glob('*/*/')
             if d.name.endswith(f'_{yaml_stem}') and _has_artifacts(d)]
    if not cands and re.match(r'^\w+_radarpillar$', yaml_stem):
        ds = yaml_stem.split('_')[0]
        cands = [d for d in runs_root.glob('*/*/')
                 if d.name.endswith(f'_{ds}') and _has_artifacts(d)]
    return max(cands, key=lambda p: p.stat().st_mtime) if cands else None


def _resbag_flops(cfg_path: Path):
    """fallback: 复用 resbag 的 thop 实测（bs=4 口径）；固定 seed 消随机抖动。"""
    import torch
    torch.manual_seed(0)
    rb = ROOT / '.claude' / 'skills' / 'resbag' / 'resbag.py'
    spec = importlib.util.spec_from_file_location('resbag', rb)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod._compute_params_flops(cfg_path, 4)


def derive_numbers(run_dir: Path | None, cfg_path: Path) -> dict:
    out = {'FLOPs': '', 'PARAMs': ''}
    p = f = None
    store = run_dir / 'model_store.yaml' if run_dir else None
    if store and store.exists():
        import yaml
        d = yaml.safe_load(store.read_text())
        p, f = d.get('params_m'), d.get('flops_g')
    if p is None:  # 无 store 或 store 缺 -> 现算
        try:
            p, _, f = _resbag_flops(cfg_path)
        except Exception as e:
            _log(f'params/FLOPs 计算失败: {type(e).__name__}', 'WARN')
    if f is not None:
        out['FLOPs'] = f'{float(f):.3f}G'
    if p is not None:
        out['PARAMs'] = f'{float(p):.3f}M'
    return out


def derive_maps(run_dir: Path | None) -> dict:
    """eval results.json -> 2D(bev)/3D mAP 列。缺 eval -> 全空。
    口径: moderate R40；mAP=在训类别均值(OB 不计入)。
    """
    out = {c: '' for c in PERF_COLS}
    if not run_dir:
        return out
    rjs = sorted((run_dir / 'eval').rglob('results.json'),
                 key=lambda p: p.stat().st_mtime) \
        if (run_dir / 'eval').exists() else []
    if not rjs:
        return out
    ret = json.loads(rjs[-1].read_text()).get('ret_dict', {})
    for k, v in ret.items():
        mm = re.match(r'(\w+?)_(3d|bev)/moderate_R40$', k)
        if not mm:
            continue
        cls, dim = mm.group(1), '3D' if mm.group(2) == '3d' else '2D'
        col = CLASS2COL.get(cls) or ('OB' if cls.upper().startswith('OB') else None)
        if col:
            out[f'{dim}{col}mAP'] = round(float(v), 2)
    for dim in ('2D', '3D'):
        vals = [out[f'{dim}{c}mAP'] for c in METRIC_CLASSES]
        if vals and all(isinstance(x, float) for x in vals):
            out[f'{dim}mAP'] = round(sum(vals) / len(vals), 2)
    return out


# -- xlsx 写入（模板感知：只写值）----------------------------------------------
def write_xlsx(xlsx: Path, rows: list[dict], stage: str):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb['MODEL'] if 'MODEL' in wb.sheetnames else wb.active

    # 列位置从表头动态读取（空列分隔列天然被排除）
    hm = {str(c.value).strip(): c.column for c in ws[1]
          if isinstance(c.value, str) and c.value.strip()}
    cols = (ARCH_COLS if stage in ('arch', 'all') else []) + \
           (PERF_COLS if stage in ('perf', 'all') else [])
    missing = [c for c in cols if c not in hm]
    if missing:
        _log(f'表头缺列: {missing}——不改表头，请人工按模板补列后重跑', 'ERROR')
        sys.exit(1)

    # 现有行索引（按 MODEL_TAG 列）
    tag_col = hm['MODEL_TAG']
    row_of = {}
    for r in range(2, 500):
        v = ws.cell(row=r, column=tag_col).value
        if v is None:
            break
        row_of[str(v).strip()] = r
    next_row = (max(row_of.values()) + 1) if row_of else 2

    # MergedCell 防护: 仅解除将被写入的单元格所在的合并区(最小干预, 其余格式不动)
    tgt_rows = list(range(2, next_row + len(rows) + 1))
    tgt_cols = set(hm.values())
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row in tgt_rows and (rng.min_col in tgt_cols or rng.max_col in tgt_cols):
            ws.unmerge_cells(str(rng))

    skipped = []  # arch 防覆写跳过的列名(仅统计)

    for row in rows:
        tag = row['MODEL_TAG']
        r = row_of.get(tag)
        is_new = r is None
        if is_new:  # 新增行: 逐列复制模板数据行(第2行)样式
            r = next_row
            next_row += 1
            row_of[tag] = r
            for col in hm.values():
                dst = ws.cell(row=r, column=col)
                dst._style = copy(ws.cell(row=2, column=col)._style)
        for c in cols:
            if row.get(c, '') == '':
                continue  # 空值不覆盖已有内容
            if c == 'MODEL_TAG' and not is_new:
                continue  # 已有行不覆盖 tag；新增行必须写 tag
            # arch 列防覆写: 已有非空值(如用户手填的 RepDwc+MDFEN 拆分写法)不重写,
            # 只有空单元格才落盘(YAML 原样编码无法表达人工拆分知识, 覆写即丢失)
            if stage in ('arch', 'all') and c in ARCH_COLS and not is_new:
                cur = ws.cell(row=r, column=hm[c]).value
                if cur not in (None, ''):
                    skipped.append(f'{c}')
                    continue
            cell = ws.cell(row=r, column=hm[c])
            # 只写值不动格式: 既有样式原样保留; 无样式单元格落盘前套模板数据行(第2行)同列样式
            # (防新写单元格变成工作簿默认字体/无线型, 与表内其他单元格观感不一致)
            if not cell.has_style:
                cell._style = copy(ws.cell(row=2, column=hm[c])._style)
            cell.value = row[c]

    if (xlsx.parent / f'~${xlsx.name}').exists():
        _log('检测到 Excel 锁文件(~$)，文件可能在 Excel 中打开，保存后请勿从 Excel 侧覆盖', 'WARN')
    if skipped:
        from collections import Counter
        cnt = ', '.join(f'{k}×{v}' for k, v in Counter(skipped).most_common())
        _log(f'arch 防覆写: {len(skipped)} 处已有值跳过 ({cnt})——确需重写请先手动清空对应单元格', 'WARN')
    wb.save(xlsx)
    _log(f'stage={stage} 已落盘 {len(rows)} 行 -> {xlsx}')


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--dataset', default='MC_DATASET',
                    help='experiments/<DATASET>/（YAML 与 xlsx 所在）')
    ap.add_argument('--xlsx', default=None, help='默认 experiments/<DATASET>/model.xlsx')
    ap.add_argument('--models', nargs='*', default=None,
                    help='yaml 名列表；默认 YAML/ 下全部（自动跳过非 dict 的指针文件）')
    ap.add_argument('--stage', choices=['arch', 'perf', 'all'], default='all',
                    help='arch=YAML 时落架构+FLOPs/PARAMs；perf=训后落 mAP；all=全部')
    ap.add_argument('--run_dir', default=None, help='显式指定唯一 run 目录（多模型时不适用）')
    ap.add_argument('--dry-run', action='store_true', help='只打印不写表')
    args = ap.parse_args()

    yaml_dir = ROOT / 'experiments' / args.dataset / 'YAML'
    xlsx = Path(args.xlsx) if args.xlsx else ROOT / 'experiments' / args.dataset / 'model.xlsx'
    stems = args.models or sorted(p.stem for p in yaml_dir.glob('*.yaml'))

    rows = []
    for stem in stems:
        yp = yaml_dir / f'{stem}.yaml'
        if not yp.exists():
            _log(f'yaml 缺失: {yp}', 'WARN')
            continue
        import yaml
        if not isinstance(yaml.safe_load(yp.read_text()), dict):
            continue  # 指针/残缺文件跳过
        cfg = load_cfg(yp)
        run_dir = find_run_dir(stem, args.run_dir if len(stems) == 1 else None)
        row = {'MODEL_TAG': stem}
        if args.stage in ('arch', 'all'):
            row.update(derive_arch(cfg))
            row.update(derive_numbers(run_dir, yp))
        if args.stage in ('perf', 'all'):
            if run_dir is None:
                _log(f'{stem}: 未找到 run 目录（mAP 留空）', 'WARN')
            row.update(derive_maps(run_dir))
        rows.append(row)
        _log(f'{stem}: run={run_dir.name if run_dir else "—"} '
             f'flops={row.get("FLOPs", "—")} params={row.get("PARAMs", "—")}')

    # 行序: baseline(<ds>_radarpillar) 最前，其余按名排序
    rows.sort(key=lambda r: (not r['MODEL_TAG'].endswith('_radarpillar'), r['MODEL_TAG']))

    if args.dry_run:
        show = [c for c in ARCH_COLS + PERF_COLS if c in rows[0]] if rows else []
        print('\t'.join(show))
        for r in rows:
            print('\t'.join(str(r.get(c, '')) for c in show))
        return
    write_xlsx(xlsx, rows, args.stage)


if __name__ == '__main__':
    main()
